"""Phase 3 step G: export v4 scores to docs/exports/verity_scores_v4.xlsx.

Sheet 1 "company_years": one row per (company, fiscal_year)
  slug, company_name, country, sasb_industry, is_financial, fiscal_year,
  report_type, narrative_word_count, e_density, s_density, g_density,
  generic_density, addon_density, e_score, s_score, g_score, composite_score,
  addon_rank, taxonomy_hash, pipeline_version

Sheet 2 "company_averages": one row per company (average of 6 yearly scores)
  slug, company_name, country, sasb_industry, is_financial, wave,
  avg_e_score, avg_s_score, avg_g_score, avg_composite, n_years_scored

Usage:
    python scripts/export_scores_v4.py
"""
from __future__ import annotations

import sys
from pathlib import Path

BACKEND = Path(__file__).resolve().parent.parent
REPO_ROOT = BACKEND.parent
sys.path.insert(0, str(BACKEND))

from dotenv import load_dotenv  # noqa: E402

load_dotenv(BACKEND / ".env")

OUTPUT_PATH = REPO_ROOT / "docs" / "exports" / "verity_scores_v4.xlsx"


def main() -> int:
    from collections import defaultdict

    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill

    import app.models.provenance
    import app.models.score
    import app.models.taxonomy  # noqa: F401
    from app.database import SessionLocal
    from app.models.institution import Institution
    from app.models.report import Report, ReportStatus
    from app.services.pipeline import PIPELINE_VERSION

    db = SessionLocal()

    reports = (
        db.query(Report)
        .join(Institution, Institution.id == Report.institution_id)
        .filter(
            Report.status == ReportStatus.scored,
            Report.pipeline_version == PIPELINE_VERSION,
            Institution.active,
        )
        .order_by(Institution.name, Report.fiscal_year)
        .all()
    )

    if not reports:
        print("No scored reports found.")
        return 1

    inst_map = {i.id: i for i in db.query(Institution).filter(Institution.active).all()}
    print(f"Exporting {len(reports)} reports to {OUTPUT_PATH}")

    wb = openpyxl.Workbook()

    # ---- Sheet 1: company_years ----
    ws1 = wb.active
    ws1.title = "company_years"

    header1 = [
        "slug", "company_name", "country", "sasb_industry", "is_financial",
        "fiscal_year", "report_type",
        "narrative_word_count",
        "e_density", "s_density", "g_density", "generic_density", "addon_density",
        "e_score", "s_score", "g_score", "composite_score", "addon_rank",
        "taxonomy_hash", "pipeline_version",
    ]
    ws1.append(header1)
    # Bold header
    hdr_font = Font(bold=True)
    hdr_fill = PatternFill("solid", fgColor="D9E1F2")
    for cell in ws1[1]:
        cell.font = hdr_font
        cell.fill = hdr_fill
        cell.alignment = Alignment(horizontal="center")

    for r in reports:
        inst = inst_map.get(r.institution_id)
        if inst is None:
            continue
        ws1.append([
            inst.slug,
            inst.name,
            inst.country,
            inst.sasb_industry,
            inst.is_financial,
            r.fiscal_year,
            r.report_type,
            r.narrative_word_count,
            round(r.e_density, 6) if r.e_density is not None else None,
            round(r.s_density, 6) if r.s_density is not None else None,
            round(r.g_density, 6) if r.g_density is not None else None,
            round(r.generic_density, 6) if r.generic_density is not None else None,
            round(r.addon_density, 6) if r.addon_density is not None else None,
            round(r.e_score, 4) if r.e_score is not None else None,
            round(r.s_score, 4) if r.s_score is not None else None,
            round(r.g_score, 4) if r.g_score is not None else None,
            round(r.composite_score, 4) if r.composite_score is not None else None,
            round(r.addon_rank, 4) if r.addon_rank is not None else None,
            r.taxonomy_version,
            r.pipeline_version,
        ])

    # Auto-width
    for col in ws1.columns:
        max_len = max((len(str(cell.value or "")) for cell in col), default=10)
        ws1.column_dimensions[col[0].column_letter].width = min(max_len + 2, 40)

    # ---- Sheet 2: company_averages ----
    ws2 = wb.create_sheet("company_averages")

    header2 = [
        "slug", "company_name", "country", "sasb_industry", "is_financial", "wave",
        "avg_e_score", "avg_s_score", "avg_g_score", "avg_composite",
        "n_years_scored",
    ]
    ws2.append(header2)
    for cell in ws2[1]:
        cell.font = hdr_font
        cell.fill = hdr_fill
        cell.alignment = Alignment(horizontal="center")

    by_inst: dict[int, list[Report]] = defaultdict(list)
    for r in reports:
        by_inst[r.institution_id].append(r)

    for iid in sorted(by_inst.keys(), key=lambda i: inst_map[i].name if i in inst_map else ""):
        rpts = by_inst[iid]
        inst = inst_map.get(iid)
        if inst is None:
            continue
        scored = [r for r in rpts if r.composite_score is not None]
        if not scored:
            continue

        def _avg(field: str, _scored: list = scored) -> float | None:
            vals = [getattr(r, field) for r in _scored if getattr(r, field) is not None]
            return round(sum(vals) / len(vals), 4) if vals else None

        ws2.append([
            inst.slug,
            inst.name,
            inst.country,
            inst.sasb_industry,
            inst.is_financial,
            inst.wave.value if inst.wave else None,
            _avg("e_score"),
            _avg("s_score"),
            _avg("g_score"),
            _avg("composite_score"),
            len(scored),
        ])

    for col in ws2.columns:
        max_len = max((len(str(cell.value or "")) for cell in col), default=10)
        ws2.column_dimensions[col[0].column_letter].width = min(max_len + 2, 40)

    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(OUTPUT_PATH))
    print(f"Saved: {OUTPUT_PATH}")
    print(f"  Sheet 'company_years':    {ws1.max_row - 1} rows")
    print(f"  Sheet 'company_averages': {ws2.max_row - 1} rows")
    return 0


if __name__ == "__main__":
    sys.exit(main())
