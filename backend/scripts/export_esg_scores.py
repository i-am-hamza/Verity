"""Export Verity ESG scores to docs/exports/.

Produces three files:
  verity_esg_scores.xlsx  — 5-sheet formatted workbook
  verity_esg_scores.csv   — flat CSV, same columns as "Scores by year"
  verity_esg_scores_column_guide.csv — one row per column with meaning + unit

Usage:
    python scripts/export_esg_scores.py
"""
from __future__ import annotations

import csv
import sys
from collections import defaultdict
from datetime import date
from pathlib import Path

BACKEND = Path(__file__).resolve().parent.parent
REPO_ROOT = BACKEND.parent
sys.path.insert(0, str(BACKEND))

from dotenv import load_dotenv  # noqa: E402

load_dotenv(BACKEND / ".env")

EXPORTS_DIR = REPO_ROOT / "docs" / "exports"

ADDED_5_SLUGS = {
    "gulf-hotels-group",
    "jazeera-steel",
    "salam-international-investment",
    "barwa-real-estate",
    "jazeera-airways",
}

# Known data-quality artifacts (slug → (fy_or_None, tag, reason))
KNOWN_ARTIFACTS: dict[str, tuple[int | None, str, str]] = {
    "rabigh-refining-petrochemical-co": (
        2020,
        "sentence_segmentation_artifact",
        "FY2020 report layout caused over-segmentation; score unreliable for that year.",
    ),
}

INDUSTRY_TO_SECTOR: dict[str, str] = {
    "Commercial Banks": "Financials",
    "Asset Management & Custody Activities": "Financials",
    "Insurance": "Financials",
    "Consumer Finance": "Financials",
    "Investment Banking & Brokerage": "Financials",
    "Mortgage Finance": "Financials",
    "Security & Commodity Exchanges": "Financials",
    "Chemicals": "Resource Transformation",
    "Iron & Steel Producers": "Resource Transformation",
    "Metals & Mining": "Resource Transformation",
    "Containers & Packaging": "Resource Transformation",
    "Aerospace & Defense": "Resource Transformation",
    "Electrical & Electronic Equipment": "Resource Transformation",
    "Industrial Machinery & Goods": "Resource Transformation",
    "Oil & Gas – Exploration & Production": "Extractives & Minerals Processing",
    "Oil & Gas – Refining & Marketing": "Extractives & Minerals Processing",
    "Oil & Gas – Services": "Extractives & Minerals Processing",
    "Oil & Gas – Midstream": "Extractives & Minerals Processing",
    "Coal Operations": "Extractives & Minerals Processing",
    "Construction Materials": "Extractives & Minerals Processing",
    "Telecommunication Services": "Technology & Communications",
    "Software & IT Services": "Technology & Communications",
    "Semiconductors": "Technology & Communications",
    "Hardware": "Technology & Communications",
    "Internet Media & Services": "Technology & Communications",
    "Electronic Manufacturing Services & Original Design Manufacturing": "Technology & Communications",
    "Hotels & Lodging": "Services",
    "Casinos & Gaming": "Services",
    "Education": "Services",
    "Media & Entertainment": "Services",
    "Professional & Commercial Services": "Services",
    "Leisure Facilities": "Services",
    "Advertising & Marketing": "Services",
    "Restaurants": "Services",
    "Meat, Poultry & Dairy": "Food & Beverage",
    "Food Retailers & Distributors": "Food & Beverage",
    "Processed Foods": "Food & Beverage",
    "Agricultural Products": "Food & Beverage",
    "Alcoholic Beverages": "Food & Beverage",
    "Non-Alcoholic Beverages": "Food & Beverage",
    "Tobacco": "Food & Beverage",
    "Real Estate": "Infrastructure",
    "Home Builders": "Infrastructure",
    "Electric Utilities & Power Generators": "Infrastructure",
    "Engineering & Construction Services": "Infrastructure",
    "Gas Utilities & Distributors": "Infrastructure",
    "Real Estate Services": "Infrastructure",
    "Waste Management": "Infrastructure",
    "Water Utilities": "Infrastructure",
    "Multiline and Specialty Retailers & Distributors": "Consumer Goods",
    "Apparel, Accessories & Footwear": "Consumer Goods",
    "Household & Personal Products": "Consumer Goods",
    "Toys & Sporting Goods": "Consumer Goods",
    "Building Products & Furnishings": "Consumer Goods",
    "E-Commerce": "Consumer Goods",
    "Airlines": "Transportation",
    "Marine Transportation": "Transportation",
    "Auto Parts": "Transportation",
    "Automobiles": "Transportation",
    "Air Freight & Logistics": "Transportation",
    "Cruise Lines": "Transportation",
    "Rail Transportation": "Transportation",
    "Road Transportation": "Transportation",
    "Car Rental & Leasing": "Transportation",
    "Health Care Delivery": "Health Care",
    "Biotechnology & Pharmaceuticals": "Health Care",
    "Drug Retailers": "Health Care",
    "Health Care Distributors": "Health Care",
    "Medical Equipment & Supplies": "Health Care",
    "Managed Care": "Health Care",
    "Biofuels": "Renewable Resources & Alternative Energy",
    "Forestry Management": "Renewable Resources & Alternative Energy",
    "Fuel Cells & Industrial Batteries": "Renewable Resources & Alternative Energy",
    "Pulp & Paper Products": "Renewable Resources & Alternative Energy",
    "Solar Technology & Project Developers": "Renewable Resources & Alternative Energy",
    "Wind Technology & Project Developers": "Renewable Resources & Alternative Energy",
}

COLUMN_GUIDE = [
    ("Company", "Company name", ""),
    ("Ticker", "Stock ticker symbol", ""),
    ("Country", "Country of primary listing", ""),
    ("SASB sector", "SASB macro-sector (11 categories)", ""),
    ("SASB industry", "SASB industry (77 categories)", ""),
    ("Financial", "Yes = one of the 20 financial companies; No = non-financial", ""),
    ("Fiscal year", "Fiscal year of the annual report", "YYYY"),
    ("Report type", "annual or integrated", ""),
    ("Environmental (0-100)", "Rank-normalised Environmental pillar score", "0–100"),
    ("Social (0-100)", "Rank-normalised Social pillar score", "0–100"),
    ("Governance (0-100)", "Rank-normalised Governance pillar score", "0–100"),
    ("General (0-100)", "Rank-normalised General ESG component score (8 terms)", "0–100"),
    ("Composite (0-100)", "Mean of E, S, G and General scores", "0–100"),
    ("Composite 3-pillar (0-100)", "Robustness: mean of E, S, G only (no General component)", "0–100"),
    ("Composite no-new3 (0-100)", "Robustness: General from 5 original terms only (no innovation/feedback/digital services)", "0–100"),
    ("E density", "Weighted Environmental term matches per 1,000 narrative words", "per 1,000 words"),
    ("S density", "Weighted Social term matches per 1,000 narrative words", "per 1,000 words"),
    ("G density", "Weighted Governance term matches per 1,000 narrative words", "per 1,000 words"),
    ("General ESG density", "General ESG terms (ESG, sustainability, GRI, innovation, feedback, digital services, …) per 1,000 words", "per 1,000 words"),
    ("Banking add-on density", "Financial-sector add-on terms per 1,000 words (financial companies only)", "per 1,000 words"),
    ("Banking add-on score", "Rank-normalised banking add-on score (financial companies only)", "0–100"),
    ("Words analysed", "Latin-script words on narrative pages used as scoring denominator", "words"),
    ("Total words", "Total Latin-script words in the full report", "words"),
    ("Share of text analysed", "Words analysed / Total words", "0–1"),
    ("Data-quality flag", "Non-empty if the report has a known quality issue", ""),
]


def _spearman(a: list[float], b: list[float]) -> float:
    n = len(a)
    if n < 2:
        return float("nan")
    ra = _rank_list(a)
    rb = _rank_list(b)
    d2 = sum((ra[i] - rb[i]) ** 2 for i in range(n))
    return round(1 - 6 * d2 / (n * (n * n - 1)), 4)


def _rank_list(vals: list[float]) -> list[float]:
    n = len(vals)
    indexed = sorted(range(n), key=lambda i: vals[i])
    ranks = [0.0] * n
    pos = 0
    while pos < n:
        end = pos + 1
        while end < n and vals[indexed[end]] == vals[indexed[pos]]:
            end += 1
        avg = (pos + end + 1) / 2.0
        for k in range(pos, end):
            ranks[indexed[k]] = avg
        pos = end
    return ranks


def _data_quality_flag(slug: str, fy: int, processing_review_status: str,
                       processing_review_reason: str | None) -> str:
    parts: list[str] = []
    art = KNOWN_ARTIFACTS.get(slug)
    if art:
        art_fy, tag, _ = art
        if art_fy is None or art_fy == fy:
            parts.append(tag)
    if processing_review_status not in ("auto_ok",):
        if processing_review_reason:
            parts.append(processing_review_reason)
        else:
            parts.append(processing_review_status)
    return "; ".join(parts)


def main() -> int:
    import openpyxl
    from openpyxl.formatting.rule import ColorScaleRule
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    import app.models.provenance  # noqa: F401
    import app.models.score  # noqa: F401
    import app.models.taxonomy  # noqa: F401
    from app.database import SessionLocal
    from app.models.institution import Institution
    from app.models.report import Report, ReportStatus
    from app.models.score import MatchEvidence
    from app.models.taxonomy import TaxonomyVersion
    from app.services.pipeline import PIPELINE_VERSION
    from app.services.sasb_materiality import DEFAULT_WEIGHT, _norm, material_categories_for
    from app.models.taxonomy import Term
    from scripts.compute_ranks import _rank_scores

    db = SessionLocal()

    all_reports = (
        db.query(Report)
        .join(Institution, Institution.id == Report.institution_id)
        .filter(
            Report.status == ReportStatus.scored,
            Report.pipeline_version == PIPELINE_VERSION,
            Institution.active,
        )
        .order_by(Institution.name, Report.fiscal_year, Report.id)
        .all()
    )

    # Deduplicate: --force can create multiple rows for the same source document.
    # Keep the row with the highest id per (source_document_id, pipeline_version).
    seen_sd: dict[tuple, int] = {}
    reports_dedup: list = []
    for r in all_reports:
        key = (r.source_document_id, PIPELINE_VERSION) if r.source_document_id else (r.institution_id, r.fiscal_year, PIPELINE_VERSION)
        if key in seen_sd:
            reports_dedup[seen_sd[key]] = r
        else:
            seen_sd[key] = len(reports_dedup)
            reports_dedup.append(r)
    reports = reports_dedup

    if not reports:
        print("No scored reports found. Run pipeline + compute_ranks.py first.")
        return 1

    tv = db.query(TaxonomyVersion).order_by(TaxonomyVersion.id.desc()).first()
    tax_hash = tv.hash if tv else "unknown"
    inst_map = {i.id: i for i in db.query(Institution).filter(Institution.active).all()}
    print(f"Exporting {len(reports)} scored reports  pipeline={PIPELINE_VERSION}  taxonomy={tax_hash[:8]}")

    # ------------------------------------------------------------------ #
    # Robustness computations (company-level averages)
    # ------------------------------------------------------------------ #
    ranked = [r for r in reports if r.e_score is not None]
    report_ids = [r.id for r in ranked]

    terms = db.query(Term).all()
    term_sasb = {t.id: (t.sasb_category_primary, t.sasb_category_secondary) for t in terms}
    term_group = {t.id: t.group for t in terms}
    term_pillar = {t.id: t.category.pillar for t in terms}

    evidence = (
        db.query(MatchEvidence)
        .filter(
            MatchEvidence.report_id.in_(report_ids),
            MatchEvidence.pipeline_version == PIPELINE_VERSION,
        )
        .all()
    )
    ev_by_report: dict[int, list] = defaultdict(list)
    for ev in evidence:
        ev_by_report[ev.report_id].append(ev)

    def _recompute_composites(material_w: float) -> list[float | None]:
        # Vary material weighting for E/S/G only; General terms have no SASB
        # category so they always get DEFAULT_WEIGHT regardless of material_w.
        e_dens, s_dens, g_dens, gen_dens = [], [], [], []
        for r in ranked:
            inst = inst_map.get(r.institution_id)
            wc = r.narrative_word_count
            if not wc or not inst:
                e_dens.append(None); s_dens.append(None)
                g_dens.append(None); gen_dens.append(None)
                continue
            mat = material_categories_for(inst.sasb_industry)
            e_raw = s_raw = g_raw = gen_raw = 0.0
            for ev in ev_by_report.get(r.id, []):
                tid = ev.term_id
                if term_group.get(tid, "") != "core":
                    continue
                pillar = term_pillar.get(tid, "")
                p_raw, s_raw_cat = term_sasb.get(tid, (None, None))
                pn = _norm(p_raw) if p_raw else None
                sn = _norm(s_raw_cat) if s_raw_cat else None
                w = material_w if mat and ((pn and pn in mat) or (sn and sn in mat)) else DEFAULT_WEIGHT
                if pillar == "Environmental":
                    e_raw += w
                elif pillar == "Social":
                    s_raw += w
                elif pillar == "Governance":
                    g_raw += w
                elif pillar == "General":
                    # General terms have no SASB category; always DEFAULT_WEIGHT
                    gen_raw += DEFAULT_WEIGHT
            k = 1000.0 / wc
            e_dens.append(e_raw * k); s_dens.append(s_raw * k)
            g_dens.append(g_raw * k); gen_dens.append(gen_raw * k)
        e_sc = _rank_scores(e_dens)
        s_sc = _rank_scores(s_dens)
        g_sc = _rank_scores(g_dens)
        gen_sc = _rank_scores(gen_dens)
        return [
            (e + s + g + gen) / 4.0
            if e is not None and s is not None and g is not None and gen is not None
            else None
            for e, s, g, gen in zip(e_sc, s_sc, g_sc, gen_sc, strict=True)
        ]

    def _sector_adjusted_composites() -> list[float | None]:
        # Adjust each density by its SASB-sector mean, then rank-score and
        # average all four components (E, S, G, General) — same as main composite.
        sectors = [INDUSTRY_TO_SECTOR.get(
            (inst_map.get(r.institution_id) or type("", (), {"sasb_industry": ""})()).sasb_industry or "", "Unknown"
        ) for r in ranked]
        se: dict[str, list[float]] = defaultdict(list)
        ss: dict[str, list[float]] = defaultdict(list)
        sg: dict[str, list[float]] = defaultdict(list)
        sgen: dict[str, list[float]] = defaultdict(list)
        for r, sec in zip(ranked, sectors, strict=True):
            if r.e_density is not None:
                se[sec].append(r.e_density)
            if r.s_density is not None:
                ss[sec].append(r.s_density)
            if r.g_density is not None:
                sg[sec].append(r.g_density)
            if r.general_density is not None:
                sgen[sec].append(r.general_density)
        se_avg   = {s: sum(v)/len(v) for s, v in se.items()}
        ss_avg   = {s: sum(v)/len(v) for s, v in ss.items()}
        sg_avg   = {s: sum(v)/len(v) for s, v in sg.items()}
        sgen_avg = {s: sum(v)/len(v) for s, v in sgen.items()}
        adj_e:   list[float | None] = []
        adj_s:   list[float | None] = []
        adj_g:   list[float | None] = []
        adj_gen: list[float | None] = []
        for r, sec in zip(ranked, sectors, strict=True):
            ea  = se_avg.get(sec,   0.0)
            sa  = ss_avg.get(sec,   0.0)
            ga  = sg_avg.get(sec,   0.0)
            gna = sgen_avg.get(sec, 0.0)
            adj_e.append(  (r.e_density       - ea)  / ea  if r.e_density       is not None and ea  > 0 else None)
            adj_s.append(  (r.s_density       - sa)  / sa  if r.s_density       is not None and sa  > 0 else None)
            adj_g.append(  (r.g_density       - ga)  / ga  if r.g_density       is not None and ga  > 0 else None)
            adj_gen.append((r.general_density  - gna) / gna if r.general_density is not None and gna > 0 else None)
        e_sc   = _rank_scores(adj_e)
        s_sc   = _rank_scores(adj_s)
        g_sc   = _rank_scores(adj_g)
        gen_sc = _rank_scores(adj_gen)
        return [
            (e + s + g + gen) / 4.0
            if e is not None and s is not None and g is not None and gen is not None
            else None
            for e, s, g, gen in zip(e_sc, s_sc, g_sc, gen_sc, strict=True)
        ]

    def _without5_composites() -> list[float | None]:
        # Re-rank within the 60-company sub-sample; compare on those 60 only.
        # Uses the same four-component composite as the main score.
        keep = [(i, r) for i, r in enumerate(ranked)
                if inst_map.get(r.institution_id) and
                inst_map[r.institution_id].slug not in ADDED_5_SLUGS]
        sub = [r for _, r in keep]
        sub_e   = _rank_scores([r.e_density       for r in sub])
        sub_s   = _rank_scores([r.s_density       for r in sub])
        sub_g   = _rank_scores([r.g_density       for r in sub])
        sub_gen = _rank_scores([r.general_density  for r in sub])
        sub_c = [
            (e + s + g + gen) / 4.0
            if e is not None and s is not None and g is not None and gen is not None
            else None
            for e, s, g, gen in zip(sub_e, sub_s, sub_g, sub_gen, strict=True)
        ]
        result: list[float | None] = [None] * len(ranked)
        for (orig_i, _), c in zip(keep, sub_c, strict=True):
            result[orig_i] = c
        return result

    # Robustness: composite_3pillar and composite_general_no_new3 from DB
    comp_3pil = [r.composite_3pillar for r in ranked]
    comp_no_new3 = [r.composite_general_no_new3 for r in ranked]

    main_comp = [r.composite_score for r in ranked]
    comp_125 = _recompute_composites(1.25)
    comp_unw = _recompute_composites(1.0)
    comp_200 = _recompute_composites(2.0)
    comp_sec = _sector_adjusted_composites()
    comp_w5  = _without5_composites()

    # Company-level averages for robustness
    def _by_inst_avg(composites: list[float | None]) -> dict[int, float]:
        d: dict[int, list[float]] = defaultdict(list)
        for r, c in zip(ranked, composites, strict=True):
            if c is not None:
                d[r.institution_id].append(c)
        return {iid: sum(vs)/len(vs) for iid, vs in d.items()}

    main_avg   = _by_inst_avg(main_comp)
    avg_3pil   = _by_inst_avg(comp_3pil)
    avg_no_new3 = _by_inst_avg(comp_no_new3)
    avg_125    = _by_inst_avg(comp_125)
    avg_unw    = _by_inst_avg(comp_unw)
    avg_200    = _by_inst_avg(comp_200)
    avg_sec    = _by_inst_avg(comp_sec)
    avg_w5     = _by_inst_avg(comp_w5)

    def _rho_inst(ref: dict[int, float], alt: dict[int, float]) -> float:
        """Company-level (averaged) Spearman."""
        common = sorted(set(ref) & set(alt))
        if len(common) < 2:
            return float("nan")
        return _spearman([ref[i] for i in common], [alt[i] for i in common])

    def _rho_cy(ref_list: list[float | None], alt_list: list[float | None]) -> float:
        """Company-year Spearman over matching non-None pairs."""
        pairs = [(a, b) for a, b in zip(ref_list, alt_list, strict=True)
                 if a is not None and b is not None]
        if len(pairs) < 2:
            return float("nan")
        return _spearman([p[0] for p in pairs], [p[1] for p in pairs])

    rho_3pil    = _rho_inst(main_avg, avg_3pil)
    rho_no_new3 = _rho_inst(main_avg, avg_no_new3)
    rho_125 = _rho_inst(main_avg, avg_125)
    rho_unw = _rho_inst(main_avg, avg_unw)
    rho_200 = _rho_inst(main_avg, avg_200)
    rho_sec = _rho_inst(main_avg, avg_sec)
    rho_w5  = _rho_inst(main_avg, avg_w5)

    rho_cy_3pil    = _rho_cy(main_comp, comp_3pil)
    rho_cy_no_new3 = _rho_cy(main_comp, comp_no_new3)
    rho_cy_125     = _rho_cy(main_comp, comp_125)
    rho_cy_unw     = _rho_cy(main_comp, comp_unw)
    rho_cy_200     = _rho_cy(main_comp, comp_200)
    rho_cy_sec     = _rho_cy(main_comp, comp_sec)
    rho_cy_w5      = _rho_cy(main_comp, comp_w5)

    # ------------------------------------------------------------------ #
    # v4 (pipeline 0.4.2) comparison
    # ------------------------------------------------------------------ #
    V4_PIPELINE = "0.4.2"
    v4_raw = (
        db.query(Report)
        .join(Institution, Institution.id == Report.institution_id)
        .filter(
            Report.status == ReportStatus.scored,
            Report.pipeline_version == V4_PIPELINE,
            Institution.active,
        )
        .order_by(Institution.name, Report.fiscal_year, Report.id)
        .all()
    )
    seen_v4: dict[tuple, int] = {}
    v4_dedup: list[Report] = []
    for r in v4_raw:
        key = (r.source_document_id, V4_PIPELINE) if r.source_document_id else (r.institution_id, r.fiscal_year, V4_PIPELINE)
        if key in seen_v4:
            v4_dedup[seen_v4[key]] = r
        else:
            seen_v4[key] = len(v4_dedup)
            v4_dedup.append(r)

    v6_key_map: dict[tuple[int, int], Report] = {(r.institution_id, r.fiscal_year): r for r in ranked}
    v4_key_map: dict[tuple[int, int], Report] = {(r.institution_id, r.fiscal_year): r for r in v4_dedup}
    paired_keys = sorted(set(v6_key_map) & set(v4_key_map))
    v6_paired   = [v6_key_map[k] for k in paired_keys]
    v4_paired   = [v4_key_map[k] for k in paired_keys]

    rho_cy_v4_comp = _rho_cy([r.composite_score for r in v6_paired], [r.composite_score for r in v4_paired])
    rho_cy_v4_e    = _rho_cy([r.e_score for r in v6_paired], [r.e_score for r in v4_paired])
    rho_cy_v4_s    = _rho_cy([r.s_score for r in v6_paired], [r.s_score for r in v4_paired])
    rho_cy_v4_g    = _rho_cy([r.g_score for r in v6_paired], [r.g_score for r in v4_paired])

    def _inst_avg(reports: list[Report], field: str) -> dict[int, float]:
        d: dict[int, list[float]] = defaultdict(list)
        for r in reports:
            v = getattr(r, field, None)
            if v is not None:
                d[r.institution_id].append(v)
        return {iid: sum(vs) / len(vs) for iid, vs in d.items()}

    v4_comp_avg = _inst_avg(v4_dedup, "composite_score")
    v4_e_avg    = _inst_avg(v4_dedup, "e_score")
    v4_s_avg    = _inst_avg(v4_dedup, "s_score")
    v4_g_avg    = _inst_avg(v4_dedup, "g_score")

    rho_inst_v4_comp = _rho_inst(main_avg,                     v4_comp_avg)
    rho_inst_v4_e    = _rho_inst(_inst_avg(ranked, "e_score"),  v4_e_avg)
    rho_inst_v4_s    = _rho_inst(_inst_avg(ranked, "s_score"),  v4_s_avg)
    rho_inst_v4_g    = _rho_inst(_inst_avg(ranked, "g_score"),  v4_g_avg)

    # Top-10 rank movers vs v4
    def _rank_dict(avg_map: dict[int, float]) -> dict[int, int]:
        srtd = sorted(avg_map.items(), key=lambda x: x[1], reverse=True)
        return {iid: i + 1 for i, (iid, _) in enumerate(srtd)}

    v6_ranks = _rank_dict(main_avg)
    v4_ranks = _rank_dict(v4_comp_avg)
    common_iids = sorted(set(v6_ranks) & set(v4_ranks))
    rank_deltas = sorted(
        [(abs(v6_ranks[i] - v4_ranks[i]), v6_ranks[i], v4_ranks[i], i) for i in common_iids],
        reverse=True,
    )

    print(f"\nRobustness correlations (Spearman rho, 4-component composite):")
    print(f"  {'Variant':<42} {'company':>8} {'cy':>8}  note")
    print(f"  {'-'*42} {'-'*8} {'-'*8}")
    print(f"  {'3-pillar (E/S/G only) vs main':<42} {rho_3pil:>8.4f} {rho_cy_3pil:>8.4f}")
    print(f"  {'no-new-3 General vs main':<42} {rho_no_new3:>8.4f} {rho_cy_no_new3:>8.4f}")
    print(f"  {'unweighted (w=1.0) vs main':<42} {rho_unw:>8.4f} {rho_cy_unw:>8.4f}  mat weight only")
    print(f"  {'w=1.25 vs main':<42} {rho_125:>8.4f} {rho_cy_125:>8.4f}  mat weight only")
    print(f"  {'w=2.0 vs main':<42} {rho_200:>8.4f} {rho_cy_200:>8.4f}  mat weight only")
    print(f"  {'sector-adjusted vs main':<42} {rho_sec:>8.4f} {rho_cy_sec:>8.4f}")
    print(f"  {'without-5 companies vs main':<42} {rho_w5:>8.4f} {rho_cy_w5:>8.4f}  60 company sub-sample")

    print(f"\nv6 vs v4 (pipeline 0.4.2) Spearman rho:")
    print(f"  {'':<18} {'company':>8} {f'cy (n={len(paired_keys)})':>14}")
    print(f"  {'Composite':<18} {rho_inst_v4_comp:>8.4f} {rho_cy_v4_comp:>12.4f}")
    print(f"  {'Environmental':<18} {rho_inst_v4_e:>8.4f} {rho_cy_v4_e:>12.4f}")
    print(f"  {'Social':<18} {rho_inst_v4_s:>8.4f} {rho_cy_v4_s:>12.4f}")
    print(f"  {'Governance':<18} {rho_inst_v4_g:>8.4f} {rho_cy_v4_g:>12.4f}")

    print(f"\nTop 10 composite rank movers v4 -> v6 (company-level):")
    for delta, r6, r4, iid in rank_deltas[:10]:
        name = inst_map[iid].name if iid in inst_map else str(iid)
        direction = "up" if r6 < r4 else "down"
        print(f"  {name:<45} v4={r4:>3d} v6={r6:>3d}  ({direction} {delta})")

    # ------------------------------------------------------------------ #
    # Build per-report row data
    # ------------------------------------------------------------------ #
    def _row(r: Report) -> dict:
        inst = inst_map.get(r.institution_id)
        if inst is None:
            return {}
        flag = _data_quality_flag(
            inst.slug, r.fiscal_year,
            r.processing_review_status,
            r.processing_review_reason,
        )
        share = (r.narrative_word_count / r.latin_word_count
                 if r.latin_word_count else None)
        return {
            "company": inst.name,
            "ticker": inst.ticker or "",
            "country": inst.country,
            "sasb_sector": INDUSTRY_TO_SECTOR.get(inst.sasb_industry or "", ""),
            "sasb_industry": inst.sasb_industry or "",
            "financial": "Yes" if inst.is_financial else "No",
            "fiscal_year": r.fiscal_year,
            "report_type": r.report_type or "",
            "e_score": round(r.e_score, 1) if r.e_score is not None else None,
            "s_score": round(r.s_score, 1) if r.s_score is not None else None,
            "g_score": round(r.g_score, 1) if r.g_score is not None else None,
            "general_score": round(r.general_score, 1) if r.general_score is not None else None,
            "composite": round(r.composite_score, 1) if r.composite_score is not None else None,
            "composite_3pillar": round(r.composite_3pillar, 1) if r.composite_3pillar is not None else None,
            "composite_general_no_new3": round(r.composite_general_no_new3, 1) if r.composite_general_no_new3 is not None else None,
            "e_density": round(r.e_density, 3) if r.e_density is not None else None,
            "s_density": round(r.s_density, 3) if r.s_density is not None else None,
            "g_density": round(r.g_density, 3) if r.g_density is not None else None,
            "general_density": round(r.general_density, 3) if r.general_density is not None else None,
            "addon_density": round(r.addon_density, 3) if (inst.is_financial and r.addon_density is not None) else None,
            "addon_score": round(r.addon_rank, 1) if (inst.is_financial and r.addon_rank is not None) else None,
            "words_analysed": r.narrative_word_count,
            "total_words": r.latin_word_count,
            "share_analysed": round(share, 4) if share is not None else None,
            "flag": flag,
        }

    rows = [_row(r) for r in reports if _row(r)]

    # ------------------------------------------------------------------ #
    # Company averages + ranks
    # ------------------------------------------------------------------ #
    by_inst: dict[int, list[Report]] = defaultdict(list)
    for r in reports:
        by_inst[r.institution_id].append(r)

    company_rows = []
    for iid in sorted(by_inst.keys(), key=lambda i: inst_map[i].name if i in inst_map else ""):
        inst = inst_map.get(iid)
        if not inst:
            continue
        rpts = [r for r in by_inst[iid] if r.composite_score is not None]
        if not rpts:
            continue

        def _avg(field: str) -> float | None:
            vals = [getattr(r, field) for r in rpts if getattr(r, field) is not None]
            return round(sum(vals) / len(vals), 2) if vals else None

        company_rows.append({
            "company": inst.name,
            "ticker": inst.ticker or "",
            "country": inst.country,
            "sasb_sector": INDUSTRY_TO_SECTOR.get(inst.sasb_industry or "", ""),
            "sasb_industry": inst.sasb_industry or "",
            "financial": "Yes" if inst.is_financial else "No",
            "avg_e": _avg("e_score"),
            "avg_s": _avg("s_score"),
            "avg_g": _avg("g_score"),
            "avg_general": _avg("general_score"),
            "avg_composite": _avg("composite_score"),
            "n_years": len(rpts),
            "iid": iid,
        })

    # Sort by avg_composite descending and assign ranks
    company_rows.sort(key=lambda x: x["avg_composite"] or 0, reverse=True)
    for i, cr in enumerate(company_rows, start=1):
        cr["rank"] = i

    # ------------------------------------------------------------------ #
    # Robustness per company
    # ------------------------------------------------------------------ #
    rob_rows = []
    for cr in company_rows:
        iid = cr["iid"]
        rob_rows.append({
            "company": cr["company"],
            "rank": cr["rank"],
            "main_composite": round(main_avg.get(iid, float("nan")), 1),
            "composite_3pillar": round(avg_3pil.get(iid, float("nan")), 1) if iid in avg_3pil else None,
            "composite_no_new3": round(avg_no_new3.get(iid, float("nan")), 1) if iid in avg_no_new3 else None,
            "unweighted": round(avg_unw.get(iid, float("nan")), 1) if iid in avg_unw else None,
            "w_125": round(avg_125.get(iid, float("nan")), 1) if iid in avg_125 else None,
            "w_200": round(avg_200.get(iid, float("nan")), 1) if iid in avg_200 else None,
            "sector_adj": round(avg_sec.get(iid, float("nan")), 1) if iid in avg_sec else None,
            "without_5": round(avg_w5.get(iid, float("nan")), 1) if iid in avg_w5 else None,
        })

    # ------------------------------------------------------------------ #
    # Excel workbook
    # ------------------------------------------------------------------ #
    EXPORTS_DIR.mkdir(parents=True, exist_ok=True)

    wb = openpyxl.Workbook()

    HDR_FONT = Font(bold=True, color="FFFFFF")
    HDR_FILL = PatternFill("solid", fgColor="1F3864")
    HDR_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
    BAND_FILL = PatternFill("solid", fgColor="F2F5FA")

    def _style_header(ws, row: int = 1) -> None:
        for cell in ws[row]:
            cell.font = HDR_FONT
            cell.fill = HDR_FILL
            cell.alignment = HDR_ALIGN

    def _auto_width(ws) -> None:
        for col in ws.columns:
            header = str(col[0].value or "")
            max_data = max((len(str(c.value or "")) for c in col[1:]), default=0)
            ws.column_dimensions[col[0].column_letter].width = min(
                max(len(header), max_data) + 2, 40
            )

    def _band_rows(ws, start: int = 2) -> None:
        for i, row in enumerate(ws.iter_rows(min_row=start)):
            if i % 2 == 1:
                for cell in row:
                    cell.fill = BAND_FILL

    def _score_color_scale(ws, col_letter: str, start_row: int, end_row: int) -> None:
        ws.conditional_formatting.add(
            f"{col_letter}{start_row}:{col_letter}{end_row}",
            ColorScaleRule(
                start_type="num", start_value=0,   start_color="F8696B",
                mid_type="num",   mid_value=50,    mid_color="FFEB84",
                end_type="num",   end_value=100,   end_color="63BE7B",
            ),
        )

    # ---- Sheet 1: Read me ----
    ws_readme = wb.active
    ws_readme.title = "Read me"
    ws_readme.column_dimensions["A"].width = 80
    meta_lines = [
        ("Verity ESG Disclosure Scores",),
        (f"Taxonomy version: {tax_hash}",),
        (f"Pipeline version: {PIPELINE_VERSION}",),
        (f"Export date: {date.today().isoformat()}",),
        ("",),
        ("Method summary",),
        ("Verity measures how much 65 large GCC-listed companies disclose about ESG topics",),
        ("in their English annual reports from FY2020–FY2025 (390 reports total).",),
        ("Scores are based on weighted keyword matches from GRI and SASB term lists.",),
        ("Term concepts: E=24, S=22 (27 phrases), G=34 (36 phrases), General=8, Banking add-on=9.",),
        ("Terms material to a company's SASB industry are weighted 1.5; others 1.0.",),
        ("Scores are rank-normalised 0-100 across all 390 reports; 100 = most disclosure.",),
        ("Composite = average of Environmental, Social, Governance, and General scores.",),
        ("",),
        ("How to read a 0–100 score",),
        ("100 = the report with the most disclosure in this sample for that pillar.",),
        ("0   = the report with the least.",),
        ("Scores show relative position within this sample, not absolute quality.",),
        ("",),
        ("Methodology page: see the Verity web dashboard → Methodology.",),
        ("",),
        ("Sheets",),
        ("Scores by year   — one row per company-year",),
        ("Company averages — six-year averages, ranked 1–65",),
        ("Robustness       — main composite vs seven variants + correlations",),
        ("Column guide     — definition and unit of every column",),
    ]
    for line in meta_lines:
        ws_readme.append(line)
    ws_readme["A1"].font = Font(bold=True, size=14)
    ws_readme["A6"].font = Font(bold=True)
    ws_readme["A13"].font = Font(bold=True)
    ws_readme["A19"].font = Font(bold=True)
    ws_readme["A20"].font = Font(bold=True)

    # ---- Sheet 2: Scores by year ----
    ws_year = wb.create_sheet("Scores by year")
    year_headers = [
        "Company", "Ticker", "Country", "SASB sector", "SASB industry", "Financial",
        "Fiscal year", "Report type",
        "Environmental (0-100)", "Social (0-100)", "Governance (0-100)", "General (0-100)",
        "Composite (0-100)", "Composite 3-pillar (0-100)", "Composite no-new3 (0-100)",
        "E density", "S density", "G density", "General ESG density",
        "Banking add-on density", "Banking add-on score",
        "Words analysed", "Total words", "Share of text analysed",
        "Data-quality flag",
    ]
    ws_year.append(year_headers)
    _style_header(ws_year)
    ws_year.freeze_panes = "B2"
    ws_year.auto_filter.ref = f"A1:{get_column_letter(len(year_headers))}1"

    for rd in rows:
        ws_year.append([
            rd["company"], rd["ticker"], rd["country"],
            rd["sasb_sector"], rd["sasb_industry"], rd["financial"],
            rd["fiscal_year"], rd["report_type"],
            rd["e_score"], rd["s_score"], rd["g_score"], rd["general_score"],
            rd["composite"], rd["composite_3pillar"], rd["composite_general_no_new3"],
            rd["e_density"], rd["s_density"], rd["g_density"], rd["general_density"],
            rd["addon_density"], rd["addon_score"],
            rd["words_analysed"], rd["total_words"], rd["share_analysed"],
            rd["flag"],
        ])

    _auto_width(ws_year)
    _band_rows(ws_year)
    n_year = ws_year.max_row
    # Color scales on 0-100 score columns (I–O = cols 9–15)
    for col_idx in range(9, 16):
        _score_color_scale(ws_year, get_column_letter(col_idx), 2, n_year)

    # ---- Sheet 3: Company averages ----
    ws_avg = wb.create_sheet("Company averages")
    avg_headers = [
        "Rank", "Company", "Ticker", "Country", "SASB sector", "SASB industry", "Financial",
        "Avg Environmental (0-100)", "Avg Social (0-100)", "Avg Governance (0-100)",
        "Avg General (0-100)", "Avg Composite (0-100)", "Years scored",
    ]
    ws_avg.append(avg_headers)
    _style_header(ws_avg)
    ws_avg.freeze_panes = "C2"
    ws_avg.auto_filter.ref = f"A1:{get_column_letter(len(avg_headers))}1"

    for cr in company_rows:
        ws_avg.append([
            cr["rank"], cr["company"], cr["ticker"], cr["country"],
            cr["sasb_sector"], cr["sasb_industry"], cr["financial"],
            cr["avg_e"], cr["avg_s"], cr["avg_g"], cr["avg_general"],
            cr["avg_composite"], cr["n_years"],
        ])

    _auto_width(ws_avg)
    _band_rows(ws_avg)
    n_avg = ws_avg.max_row
    for col_idx in range(8, 13):  # Avg E/S/G/General/Composite (cols 8–12)
        _score_color_scale(ws_avg, get_column_letter(col_idx), 2, n_avg)

    # ---- Sheet 4: Robustness ----
    ws_rob = wb.create_sheet("Robustness")
    rob_headers = [
        "Rank", "Company",
        "Main composite (0-100)", "3-pillar composite", "No-new-3 composite",
        "Unweighted (w=1.0)", "w=1.25", "w=2.0",
        "Sector-adjusted", "Without 5 added companies",
    ]
    ws_rob.append(rob_headers)
    _style_header(ws_rob)
    ws_rob.freeze_panes = "C2"

    for rb in rob_rows:
        ws_rob.append([
            rb["rank"], rb["company"],
            rb["main_composite"], rb["composite_3pillar"], rb["composite_no_new3"],
            rb["unweighted"], rb["w_125"], rb["w_200"],
            rb["sector_adj"], rb["without_5"],
        ])

    # Correlation table below
    n_rob = ws_rob.max_row
    ws_rob.append([])
    ws_rob.append(["Spearman correlations (4-component composite: E, S, G, General each 0-100)", "", "company", "company-year"])
    ws_rob.append(["3-pillar (E/S/G only, no General) vs main",    "", round(rho_3pil, 4),    round(rho_cy_3pil, 4)])
    ws_rob.append(["No-new-3 General vs main",                     "", round(rho_no_new3, 4), round(rho_cy_no_new3, 4)])
    ws_rob.append(["Unweighted (w=1.0) vs main (mat w=1.5)",       "", round(rho_unw, 4),     round(rho_cy_unw, 4)])
    ws_rob.append(["w=1.25 vs main (mat w=1.5)",                   "", round(rho_125, 4),     round(rho_cy_125, 4)])
    ws_rob.append(["w=2.0 vs main (mat w=1.5)",                    "", round(rho_200, 4),     round(rho_cy_200, 4)])
    ws_rob.append(["Sector-adjusted vs main",                      "", round(rho_sec, 4),     round(rho_cy_sec, 4)])
    ws_rob.append(["Without 5 added companies vs main (60 co.)",   "", round(rho_w5, 4),      round(rho_cy_w5, 4)])
    ws_rob.append([])
    ws_rob.append(["v6 vs v4 (pipeline 0.4.2) Spearman rho",      "", "company", f"company-year (n={len(paired_keys)})"])
    ws_rob.append(["Composite v6 vs v4",    "", round(rho_inst_v4_comp, 4), round(rho_cy_v4_comp, 4)])
    ws_rob.append(["Environmental v6 vs v4", "", round(rho_inst_v4_e, 4),  round(rho_cy_v4_e, 4)])
    ws_rob.append(["Social v6 vs v4",        "", round(rho_inst_v4_s, 4),  round(rho_cy_v4_s, 4)])
    ws_rob.append(["Governance v6 vs v4",    "", round(rho_inst_v4_g, 4),  round(rho_cy_v4_g, 4)])

    # Bold the two section headers
    for row in ws_rob.iter_rows(min_row=n_rob + 2, max_row=n_rob + 2):
        for cell in row:
            cell.font = Font(bold=True)
    for row in ws_rob.iter_rows(min_row=n_rob + 11, max_row=n_rob + 11):
        for cell in row:
            cell.font = Font(bold=True)

    _auto_width(ws_rob)
    _band_rows(ws_rob)

    # ---- Sheet 5: Column guide ----
    ws_guide = wb.create_sheet("Column guide")
    ws_guide.append(["Column", "Meaning", "Unit"])
    _style_header(ws_guide)
    for col, meaning, unit in COLUMN_GUIDE:
        ws_guide.append([col, meaning, unit])
    _auto_width(ws_guide)

    # Save XLSX
    xlsx_path = EXPORTS_DIR / "verity_esg_scores.xlsx"
    wb.save(str(xlsx_path))
    print(f"Saved: {xlsx_path}")
    print(f"  'Scores by year':    {ws_year.max_row - 1} rows")
    print(f"  'Company averages':  {ws_avg.max_row - 1} rows")
    print(f"  'Robustness':        {n_rob - 1} company rows + correlation table")

    # ------------------------------------------------------------------ #
    # CSV: Scores by year
    # ------------------------------------------------------------------ #
    csv_headers = [
        "company", "ticker", "country", "sasb_sector", "sasb_industry", "financial",
        "fiscal_year", "report_type",
        "environmental_score", "social_score", "governance_score", "general_score",
        "composite_score", "composite_3pillar", "composite_general_no_new3",
        "e_density", "s_density", "g_density", "general_esg_density",
        "banking_addon_density", "banking_addon_score",
        "words_analysed", "total_words", "share_of_text_analysed",
        "data_quality_flag",
    ]
    csv_path = EXPORTS_DIR / "verity_esg_scores.csv"
    with csv_path.open("w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=csv_headers)
        w.writeheader()
        for rd in rows:
            w.writerow({
                "company": rd["company"],
                "ticker": rd["ticker"],
                "country": rd["country"],
                "sasb_sector": rd["sasb_sector"],
                "sasb_industry": rd["sasb_industry"],
                "financial": rd["financial"],
                "fiscal_year": rd["fiscal_year"],
                "report_type": rd["report_type"],
                "environmental_score": f"{rd['e_score']:.1f}" if rd["e_score"] is not None else "",
                "social_score": f"{rd['s_score']:.1f}" if rd["s_score"] is not None else "",
                "governance_score": f"{rd['g_score']:.1f}" if rd["g_score"] is not None else "",
                "general_score": f"{rd['general_score']:.1f}" if rd["general_score"] is not None else "",
                "composite_score": f"{rd['composite']:.1f}" if rd["composite"] is not None else "",
                "composite_3pillar": f"{rd['composite_3pillar']:.1f}" if rd["composite_3pillar"] is not None else "",
                "composite_general_no_new3": f"{rd['composite_general_no_new3']:.1f}" if rd["composite_general_no_new3"] is not None else "",
                "e_density": f"{rd['e_density']:.4f}" if rd["e_density"] is not None else "",
                "s_density": f"{rd['s_density']:.4f}" if rd["s_density"] is not None else "",
                "g_density": f"{rd['g_density']:.4f}" if rd["g_density"] is not None else "",
                "general_esg_density": f"{rd['general_density']:.4f}" if rd["general_density"] is not None else "",
                "banking_addon_density": f"{rd['addon_density']:.4f}" if rd["addon_density"] is not None else "",
                "banking_addon_score": f"{rd['addon_score']:.1f}" if rd["addon_score"] is not None else "",
                "words_analysed": rd["words_analysed"],
                "total_words": rd["total_words"],
                "share_of_text_analysed": f"{rd['share_analysed']:.4f}" if rd["share_analysed"] is not None else "",
                "data_quality_flag": rd["flag"],
            })
    print(f"Saved: {csv_path}  ({len(rows)} rows)")

    # ------------------------------------------------------------------ #
    # CSV: Column guide
    # ------------------------------------------------------------------ #
    guide_path = EXPORTS_DIR / "verity_esg_scores_column_guide.csv"
    with guide_path.open("w", newline="", encoding="utf-8") as fh:
        w2 = csv.writer(fh)
        w2.writerow(["column", "meaning", "unit"])
        for col, meaning, unit in COLUMN_GUIDE:
            w2.writerow([col, meaning, unit])
    print(f"Saved: {guide_path}")

    # Remove old file if present
    old = EXPORTS_DIR / "verity_scores_v4.xlsx"
    if old.exists():
        old.unlink()
        print(f"Removed: {old}")

    return 0


if __name__ == "__main__":
    sys.exit(main())
