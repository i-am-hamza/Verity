"""Read Verity_Report_Sourcing.xlsx 'To source' sheet and print all rows with their Result."""
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')

import openpyxl
from pathlib import Path

wb = openpyxl.load_workbook(Path(r"E:\9. Verity\docs\methodology\Verity_Report_Sourcing.xlsx"), read_only=True, data_only=True)
print("Sheets:", wb.sheetnames)

ws = wb["To source"]
rows = list(ws.iter_rows(values_only=True))

# Print header and all rows
print(f"\nHeader: {rows[0]}")
print(f"Total data rows: {len(rows)-1}\n")

for i, row in enumerate(rows[1:], 1):
    print(f"{i:3d}  {row}")

wb.close()
