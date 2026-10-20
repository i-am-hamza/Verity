"""Compute material-only ESG disclosure scores (v4, read-only).

For each of the 390 deduplicated company-years (pipeline 0.4.2,
taxonomy 313c0eabe0...) this script:
  1. Filters existing term matches to those whose SASB category is material
     for the company's SASB industry.
  2. For the 19 financial companies, also includes banking add-on terms
     where material (attributed to E/S/G pillar via ADDON_PILLAR_MAP).
  3. Computes density = material mentions / narrative_word_count * 1000,
     weight = 1.0 per match (filter replaces the 1.5× materiality bonus).
  4. Ranks 0-10 across all 390 company-years per pillar (same pool and
     formula as compute_ranks.py).
  5. ESG_mat = simple average(E_mat_score, S_mat_score, G_mat_score).
  6. Exports docs/exports/verity_material_scores.xlsx with:
       - "Scores" sheet: one row per company-year
       - "Coverage" sheet: material terms available per SASB industry/pillar
       - "Correlation" sheet: Spearman(ESG_mat, main composite)
  7. Does NOT write to the database.

Judgment calls documented in DECISIONS.md:
  - Addon terms attributed to E/S/G via ADDON_PILLAR_MAP (see below).
  - Terms with no SASB category (e.g. corporate governance, board of directors)
    are NOT included in the material-only score.
  - Generic terms (ESG, sustainability, GRI, etc.) are excluded.
"""
from __future__ import annotations

import sys
from collections import defaultdict
from pathlib import Path
from typing import Optional

BACKEND = Path(__file__).resolve().parent.parent
REPO_ROOT = BACKEND.parent
sys.path.insert(0, str(BACKEND))

from dotenv import load_dotenv  # noqa: E402

load_dotenv(BACKEND / ".env")

import app.models.institution  # noqa: F401
import app.models.provenance   # noqa: F401
import app.models.report       # noqa: F401
import app.models.review       # noqa: F401
import app.models.run          # noqa: F401
import app.models.score        # noqa: F401
import app.models.taxonomy     # noqa: F401

from app.database import SessionLocal
from app.models.institution import Institution
from app.models.report import Report, ReportStatus
from app.models.score import MatchEvidence
from app.models.taxonomy import Category, Term, TaxonomyVersion
from app.services.pipeline import PIPELINE_VERSION
from app.services.sasb_materiality import _norm, material_categories_for
from sqlalchemy import func, select

XLSX_OUT = REPO_ROOT / "docs" / "exports" / "verity_material_scores.xlsx"

# Addon terms have pillar="Addon" in the DB.  For the material-only score
# we need to attribute each addon term to an E/S/G pillar.
# Mapping is by phrase; rationale in DECISIONS.md.
ADDON_PILLAR_MAP: dict[str, str] = {
    "green financing":       "Environmental",
    "sustainable finance":   "Environmental",
    "sustainable financing": "Environmental",
    "financial inclusion":   "Social",
    "fair lending":          "Social",
    "financial literacy":    "Social",
    "responsible lending":   "Social",
    "anti-money laundering": "Governance",
    "stress testing":        "Governance",
}

ESG_PILLARS = ("Environmental", "Social", "Governance")


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _rank_scores(values: list[Optional[float]]) -> list[Optional[float]]:
    """0-10 ascending average rank, identical to compute_ranks.py."""
    indexed = [(v, i) for i, v in enumerate(values) if v is not None]
    if not indexed:
        return [None] * len(values)
    n = len(indexed)
    result: list[Optional[float]] = [None] * len(values)
    if n == 1:
        result[indexed[0][1]] = 0.0
        return result
    sorted_vals = sorted(indexed, key=lambda x: x[0])
    avg_ranks: dict[int, float] = {}
    pos = 0
    while pos < n:
        run_end = pos + 1
        while run_end < n and sorted_vals[run_end][0] == sorted_vals[pos][0]:
            run_end += 1
        avg_rank = (pos + run_end + 1) / 2.0
        for k in range(pos, run_end):
            avg_ranks[sorted_vals[k][1]] = avg_rank
        pos = run_end
    for _v, i in indexed:
        result[i] = round(10.0 * (avg_ranks[i] - 1) / (n - 1), 4)
    return result


def _spearman(x: list[Optional[float]], y: list[Optional[float]]) -> float:
    """Spearman correlation (Pearson on ranks) with average-rank tie handling."""
    pairs = [(a, b) for a, b in zip(x, y) if a is not None and b is not None]
    if len(pairs) < 3:
        return float("nan")
    xs, ys = zip(*pairs)

    def _ranks(vals: tuple[float, ...]) -> list[float]:
        n = len(vals)
        order = sorted(range(n), key=lambda i: vals[i])
        r = [0.0] * n
        pos = 0
        while pos < n:
            end = pos + 1
            while end < n and vals[order[end]] == vals[order[pos]]:
                end += 1
            avg = (pos + end + 1) / 2.0
            for k in range(pos, end):
                r[order[k]] = avg
            pos = end
        return r

    rx = _ranks(xs)
    ry = _ranks(ys)
    n = len(xs)
    mx = sum(rx) / n
    my = sum(ry) / n
    num = sum((a - mx) * (b - my) for a, b in zip(rx, ry))
    den = (sum((a - mx) ** 2 for a in rx) * sum((b - my) ** 2 for b in ry)) ** 0.5
    return num / den if den else float("nan")


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def main() -> int:  # noqa: C901 (complex but linear)
    db = SessionLocal()

    # ------------------------------------------------------------------
    # 1. Load taxonomy version ID for 313c0eabe0...
    # ------------------------------------------------------------------
    tv = db.query(TaxonomyVersion).filter(
        TaxonomyVersion.hash.like("313c0eabe0%")
    ).one()
    tv_id = tv.id
    print(f"Taxonomy version: id={tv_id}  hash={tv.hash[:16]}…  ({tv.note})")

    # ------------------------------------------------------------------
    # 2. Load terms: id → (pillar_for_material, group, sasb_cat_primary, sasb_cat_secondary)
    # ------------------------------------------------------------------
    terms_rows = (
        db.query(Term, Category.pillar)
        .join(Category, Category.id == Term.category_id)
        .all()
    )
    # term_info: term_id → dict
    term_info: dict[int, dict] = {}
    for term, pillar in terms_rows:
        # Resolve pillar for addon terms
        if term.group == "addon":
            effective_pillar = ADDON_PILLAR_MAP.get(term.phrase)
        elif term.group == "core":
            effective_pillar = pillar  # Environmental / Social / Governance
        else:
            effective_pillar = None  # generic → skip
        term_info[term.id] = {
            "phrase":   term.phrase,
            "group":    term.group,
            "pillar":   effective_pillar,
            "scat1":    _norm(term.sasb_category_primary) if term.sasb_category_primary else None,
            "scat2":    _norm(term.sasb_category_secondary) if term.sasb_category_secondary else None,
        }
    print(f"Terms loaded: {len(term_info)}")

    # ------------------------------------------------------------------
    # 3. Load and deduplicate scored reports (same logic as compute_ranks.py)
    # ------------------------------------------------------------------
    all_scored = (
        db.query(Report)
        .join(Institution, Institution.id == Report.institution_id)
        .filter(
            Report.status == ReportStatus.scored,
            Report.pipeline_version == PIPELINE_VERSION,
            Institution.active == True,  # noqa: E712
        )
        .order_by(Institution.id, Report.fiscal_year, Report.id)
        .all()
    )
    seen: dict[tuple, int] = {}
    reports_list: list[Report] = []
    for r in all_scored:
        key = (
            (r.source_document_id, PIPELINE_VERSION)
            if r.source_document_id
            else (r.institution_id, r.fiscal_year, PIPELINE_VERSION)
        )
        if key in seen:
            reports_list[seen[key]] = r
        else:
            seen[key] = len(reports_list)
            reports_list.append(r)
    print(f"Reports (dedup): {len(reports_list)}  (raw: {len(all_scored)})")

    report_ids = [r.id for r in reports_list]

    # ------------------------------------------------------------------
    # 4. Load institution data
    # ------------------------------------------------------------------
    inst_map: dict[int, Institution] = {
        i.id: i
        for i in db.query(Institution).filter(Institution.active == True).all()  # noqa: E712
    }

    # ------------------------------------------------------------------
    # 5. Load aggregated match evidence: {report_id: {term_id: count}}
    # ------------------------------------------------------------------
    print("Loading match evidence…")
    me_rows = (
        db.query(
            MatchEvidence.report_id,
            MatchEvidence.term_id,
            func.count(MatchEvidence.id).label("cnt"),
        )
        .filter(
            MatchEvidence.pipeline_version == PIPELINE_VERSION,
            MatchEvidence.taxonomy_version_id == tv_id,
            MatchEvidence.report_id.in_(report_ids),
        )
        .group_by(MatchEvidence.report_id, MatchEvidence.term_id)
        .all()
    )
    matches_by_report: dict[int, dict[int, int]] = defaultdict(dict)
    for row in me_rows:
        matches_by_report[row.report_id][row.term_id] = row.cnt
    print(f"  Match evidence rows (aggregated): {len(me_rows)}")

    db.close()

    # ------------------------------------------------------------------
    # 6. Compute per-company-year material densities
    # ------------------------------------------------------------------
    rows_out: list[dict] = []

    for report in reports_list:
        inst = inst_map.get(report.institution_id)
        if inst is None:
            continue

        mat_cats = material_categories_for(inst.sasb_industry)
        n_words = report.narrative_word_count or 0

        # Accumulate material match counts per ESG pillar
        pillar_counts: dict[str, float] = {"Environmental": 0.0, "Social": 0.0, "Governance": 0.0}
        report_matches = matches_by_report.get(report.id, {})

        for term_id, count in report_matches.items():
            ti = term_info.get(term_id)
            if ti is None:
                continue
            if ti["group"] == "generic":
                continue
            if ti["group"] == "addon" and not inst.is_financial:
                continue
            pillar = ti["pillar"]
            if pillar not in ESG_PILLARS:
                continue
            # Materiality check: at least one SASB category is material
            scat1, scat2 = ti["scat1"], ti["scat2"]
            if not (scat1 or scat2):
                continue  # no SASB category → not material for any industry
            is_mat = (scat1 and scat1 in mat_cats) or (scat2 and scat2 in mat_cats)
            if not is_mat:
                continue
            pillar_counts[pillar] += count  # weight = 1.0 per match

        # Density (0 if no narrative words)
        if n_words > 0:
            e_d = pillar_counts["Environmental"] / n_words * 1000
            s_d = pillar_counts["Social"]        / n_words * 1000
            g_d = pillar_counts["Governance"]    / n_words * 1000
        else:
            e_d = s_d = g_d = 0.0

        rows_out.append({
            "report_id":    report.id,
            "inst_id":      inst.id,
            "slug":         inst.slug,
            "company":      inst.name,
            "ticker":       inst.ticker or "",
            "country":      inst.country,
            "sasb_industry": inst.sasb_industry or "",
            "financial":    "Yes" if inst.is_financial else "No",
            "fiscal_year":  report.fiscal_year,
            "report_type":  report.report_type or "",
            # material densities (raw)
            "e_mat_density": round(e_d, 6),
            "s_mat_density": round(s_d, 6),
            "g_mat_density": round(g_d, 6),
            # main scores (0-10 from compute_ranks.py)
            "e_score":        report.e_score,
            "s_score":        report.s_score,
            "g_score":        report.g_score,
            "composite_score": report.composite_score,
        })

    rows_out.sort(key=lambda r: (r["slug"], r["fiscal_year"]))
    print(f"Company-years computed: {len(rows_out)}")

    # ------------------------------------------------------------------
    # 7. Rank material densities 0-10 (same pool and formula as main score)
    # ------------------------------------------------------------------
    e_densities = [r["e_mat_density"] for r in rows_out]
    s_densities = [r["s_mat_density"] for r in rows_out]
    g_densities = [r["g_mat_density"] for r in rows_out]

    e_mat_scores = _rank_scores(e_densities)
    s_mat_scores = _rank_scores(s_densities)
    g_mat_scores = _rank_scores(g_densities)

    for i, row in enumerate(rows_out):
        e_m = e_mat_scores[i]
        s_m = s_mat_scores[i]
        g_m = g_mat_scores[i]
        row["E_mat"] = round(e_m, 4) if e_m is not None else None
        row["S_mat"] = round(s_m, 4) if s_m is not None else None
        row["G_mat"] = round(g_m, 4) if g_m is not None else None
        if e_m is not None and s_m is not None and g_m is not None:
            row["ESG_mat"] = round((e_m + s_m + g_m) / 3.0, 4)
        else:
            row["ESG_mat"] = None

    # ------------------------------------------------------------------
    # 8. Coverage report: material term counts per SASB industry × pillar
    # ------------------------------------------------------------------
    # Build material term availability: for each SASB industry, which terms are material?
    all_sasb_industries = sorted({r["sasb_industry"] for r in rows_out if r["sasb_industry"]})

    coverage_rows: list[dict] = []
    for industry in all_sasb_industries:
        mat_cats = material_categories_for(industry)
        # Is this a financial industry? (any company in our universe is_financial?)
        is_fin_industry = any(
            r["financial"] == "Yes"
            for r in rows_out if r["sasb_industry"] == industry
        )
        counts: dict[str, int] = {"Environmental": 0, "Social": 0, "Governance": 0}
        for ti in term_info.values():
            if ti["group"] == "generic":
                continue
            if ti["group"] == "addon" and not is_fin_industry:
                continue
            pillar = ti["pillar"]
            if pillar not in ESG_PILLARS:
                continue
            scat1, scat2 = ti["scat1"], ti["scat2"]
            if not (scat1 or scat2):
                continue
            if (scat1 and scat1 in mat_cats) or (scat2 and scat2 in mat_cats):
                counts[pillar] += 1

        # Company-years in this industry
        cy_rows = [r for r in rows_out if r["sasb_industry"] == industry]
        n_cy = len(cy_rows)
        zero_e = sum(1 for r in cy_rows if r["e_mat_density"] == 0.0)
        zero_s = sum(1 for r in cy_rows if r["s_mat_density"] == 0.0)
        zero_g = sum(1 for r in cy_rows if r["g_mat_density"] == 0.0)
        coverage_rows.append({
            "sasb_industry":       industry,
            "financial_industry":  "Yes" if is_fin_industry else "No",
            "material_terms_E":    counts["Environmental"],
            "material_terms_S":    counts["Social"],
            "material_terms_G":    counts["Governance"],
            "n_company_years":     n_cy,
            "zero_E_score":        zero_e,
            "zero_S_score":        zero_s,
            "zero_G_score":        zero_g,
        })

    # ------------------------------------------------------------------
    # 9. Spearman correlations
    # ------------------------------------------------------------------
    # Company-year level
    esg_mat_cy   = [r["ESG_mat"]        for r in rows_out]
    composite_cy = [r["composite_score"] for r in rows_out]
    r_cy = _spearman(esg_mat_cy, composite_cy)

    # Company (institution) level: mean across scored fiscal years
    by_inst: dict[int, list[tuple[float, float]]] = defaultdict(list)
    for r in rows_out:
        if r["ESG_mat"] is not None and r["composite_score"] is not None:
            by_inst[r["inst_id"]].append((r["ESG_mat"], r["composite_score"]))
    inst_esg_mat   = [sum(v for v, _ in vals) / len(vals) for vals in by_inst.values()]
    inst_composite = [sum(v for _, v in vals) / len(vals) for vals in by_inst.values()]
    r_inst = _spearman(inst_esg_mat, inst_composite)

    print(f"\nSpearman correlation:")
    print(f"  Company-year level:  r = {r_cy:.4f}  (n={sum(1 for a,b in zip(esg_mat_cy, composite_cy) if a is not None and b is not None)})")
    print(f"  Company level:       r = {r_inst:.4f}  (n={len(inst_esg_mat)})")

    # Print basic statistics
    valid_esg_mat = [r["ESG_mat"] for r in rows_out if r["ESG_mat"] is not None]
    print(f"\nMaterial score summary (company-years with score):")
    print(f"  ESG_mat: mean={sum(valid_esg_mat)/len(valid_esg_mat):.2f}  "
          f"min={min(valid_esg_mat):.2f}  max={max(valid_esg_mat):.2f}  "
          f"n={len(valid_esg_mat)}")

    # ------------------------------------------------------------------
    # 10. Export XLSX
    # ------------------------------------------------------------------
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = openpyxl.Workbook()
    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(bold=True, color="FFFFFF")

    def _write_header(ws: object, cols: list[str]) -> None:
        for ci, col in enumerate(cols, start=1):
            cell = ws.cell(row=1, column=ci, value=col)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")
        ws.freeze_panes = "A2"

    # ---- Sheet 1: Scores -----------------------------------------------
    ws_scores = wb.active
    ws_scores.title = "Scores"
    score_cols = [
        "company", "ticker", "country", "sasb_industry", "financial",
        "fiscal_year", "report_type",
        "E_mat", "S_mat", "G_mat", "ESG_mat",
        "e_mat_density", "s_mat_density", "g_mat_density",
        "e_score", "s_score", "g_score", "composite_score",
    ]
    _write_header(ws_scores, score_cols)
    for ri, row in enumerate(rows_out, start=2):
        for ci, col in enumerate(score_cols, start=1):
            ws_scores.cell(row=ri, column=ci, value=row.get(col))

    # Column widths
    col_widths = {"company": 50, "ticker": 10, "country": 25,
                  "sasb_industry": 45, "financial": 10, "fiscal_year": 12,
                  "report_type": 12}
    for ci, col in enumerate(score_cols, start=1):
        w = col_widths.get(col, 14)
        ws_scores.column_dimensions[openpyxl.utils.get_column_letter(ci)].width = w

    # ---- Sheet 2: Coverage ---------------------------------------------
    ws_cov = wb.create_sheet("Coverage")
    cov_cols = [
        "sasb_industry", "financial_industry",
        "material_terms_E", "material_terms_S", "material_terms_G",
        "n_company_years",
        "zero_E_score", "zero_S_score", "zero_G_score",
    ]
    _write_header(ws_cov, cov_cols)
    for ri, row in enumerate(coverage_rows, start=2):
        for ci, col in enumerate(cov_cols, start=1):
            ws_cov.cell(row=ri, column=ci, value=row.get(col))
    ws_cov.column_dimensions["A"].width = 50
    ws_cov.column_dimensions["B"].width = 18
    for ci in range(3, len(cov_cols) + 1):
        ws_cov.column_dimensions[openpyxl.utils.get_column_letter(ci)].width = 18

    # ---- Sheet 3: Correlation ------------------------------------------
    ws_corr = wb.create_sheet("Correlation")
    corr_data = [
        ["Level", "n", "Spearman r", "Interpretation"],
        ["Company-year",
         sum(1 for a, b in zip(esg_mat_cy, composite_cy) if a is not None and b is not None),
         round(r_cy, 4),
         "ESG_mat vs composite_score per company-year"],
        ["Company (mean across FYs)",
         len(inst_esg_mat),
         round(r_inst, 4),
         "mean(ESG_mat) vs mean(composite_score) per institution"],
    ]
    for ri, row_data in enumerate(corr_data, start=1):
        for ci, val in enumerate(row_data, start=1):
            cell = ws_corr.cell(row=ri, column=ci, value=val)
            if ri == 1:
                cell.fill = header_fill
                cell.font = header_font
    ws_corr.column_dimensions["A"].width = 30
    ws_corr.column_dimensions["B"].width = 10
    ws_corr.column_dimensions["C"].width = 14
    ws_corr.column_dimensions["D"].width = 55

    wb.save(XLSX_OUT)
    print(f"\nXLSX -> {XLSX_OUT}")

    # ------------------------------------------------------------------
    # 11. Print coverage summary
    # ------------------------------------------------------------------
    print("\n--- Coverage summary ---")
    print(f"{'SASB Industry':<50} {'Fin':>3}  {'E':>3} {'S':>3} {'G':>3}  {'N':>4}  {'0E':>4} {'0S':>4} {'0G':>4}")
    print("-" * 90)
    for c in coverage_rows:
        print(
            f"{c['sasb_industry']:<50} {c['financial_industry']:>3}  "
            f"{c['material_terms_E']:>3} {c['material_terms_S']:>3} {c['material_terms_G']:>3}  "
            f"{c['n_company_years']:>4}  "
            f"{c['zero_E_score']:>4} {c['zero_S_score']:>4} {c['zero_G_score']:>4}"
        )

    return 0


if __name__ == "__main__":
    sys.exit(main())
