"""
export_review_xlsx.py  --  Build docs/exports/financials_review.xlsx

Two tabs:
  "Must check"      -- BS_MISMATCH, NP_ATTR_GT_TOTAL, LEVERAGE_NEGATIVE_EQUITY,
                       plus NBK and Bank Muscat ROE/EPS rows (ratio-changing flags)
  "For information" -- EPS_MISMATCH, ROE_NPA_FALLBACK, OCR_HEAVY
                       (flagged but not ratio-changing)

Pure FS_BOUNDARY_NOT_FOUND rows (no FS in the PDF) are excluded from both tabs;
they appear in docs/exports/fs_sourcing_checklist.xlsx instead.
"""
from __future__ import annotations

import csv
import re
import sys
from pathlib import Path

BACKEND = Path(__file__).resolve().parent.parent
REPO_ROOT = BACKEND.parent
sys.path.insert(0, str(BACKEND))

# ---- config -----------------------------------------------------------------

CSV_IN   = REPO_ROOT / "docs" / "exports" / "financials_extracted.csv"
XLSX_OUT = REPO_ROOT / "docs" / "exports" / "financials_review.xlsx"

_FS_ONLY_RE = re.compile(r"^FS_BOUNDARY_NOT_FOUND[^;]*$")

_EV_FIELDS = [
    "total_assets",
    "total_liabilities",
    "total_equity",
    "equity_attributable",
    "revenue",
    "gross_income_bank",
    "net_profit",
    "net_profit_attributable",
    "eps_numerator",
    "basic_eps",
    "shares_outstanding",
    "shares_eps_note",
    # Debt components
    "st_borrowings",
    "lt_borrowings",
    "sukuk_bonds_notes",
    "lease_liabilities_nb",
    "due_to_banks_fi",
    "sukuk_issued_bank",
    "subordinated_debt_bank",
    "total_borrowings",
]

_RATIO_FIELDS = [
    "total_debt", "total_debt_status",
    "leverage", "leverage_alt",
    "roa", "roe", "net_profit_margin", "net_margin_alt",
    "equity_ordinary",
    "eps_in_fils_baisa",
]

_FLAG_PREFIXES = [
    "BS_MISMATCH",
    "NP_ATTR_GT_TOTAL",
    "EPS_MISMATCH",
    "ROE_NPA_FALLBACK",
    "LEVERAGE_NEGATIVE_EQUITY",
    "FS_BOUNDARY_NOT_FOUND",
    "OCR_HEAVY",
]

# ---- tab-routing constants --------------------------------------------------

_MUST_CHECK_FLAGS  = {"BS_MISMATCH", "NP_ATTR_GT_TOTAL", "LEVERAGE_NEGATIVE_EQUITY"}
_FOR_INFO_FLAGS    = {"EPS_MISMATCH", "ROE_NPA_FALLBACK", "OCR_HEAVY"}
# NBK and Bank Muscat EPS/ROE fallback rows are elevated to Must check because
# their EPS numerator extraction is known to be unreliable (currency unit issues).
_NBK_MUSCAT_SLUGS  = {"national-bank-of-kuwait", "bank-muscat-bkmb"}
_NBK_MUSCAT_EPS    = {"EPS_MISMATCH", "ROE_NPA_FALLBACK"}


def _flag_types(flags_str: str) -> str:
    """Return comma-separated list of flag type codes present in flags_str."""
    if not flags_str:
        return ""
    types = []
    for f in flags_str.split(";"):
        f = f.strip()
        for p in _FLAG_PREFIXES:
            if f.startswith(p):
                if p not in types:
                    types.append(p)
                break
    return ", ".join(types)


_OTHER_FLAG_PREFIXES = {
    "BS_MISMATCH", "NP_ATTR_GT_TOTAL", "EPS_MISMATCH",
    "ROE_NPA_FALLBACK", "LEVERAGE_NEGATIVE_EQUITY", "OCR_HEAVY",
}


def _is_true_no_fs(flags_str: str) -> bool:
    """True when FS_BOUNDARY_NOT_FOUND is present and no other flag type is.

    Cannot split on ";" because the FS_BOUNDARY flag message itself contains
    a semicolon ("detected; PDF may contain…").  Check by substring instead.
    """
    if not flags_str:
        return False
    if "FS_BOUNDARY_NOT_FOUND" not in flags_str:
        return False
    return not any(p in flags_str for p in _OTHER_FLAG_PREFIXES)


def _has_any_flag(flags_str: str, flag_set: set[str]) -> bool:
    if not flags_str:
        return False
    for f in flags_str.split(";"):
        f = f.strip()
        for p in flag_set:
            if f.startswith(p):
                return True
    return False


def _is_must_check(row: dict[str, str]) -> bool:
    flags = row.get("flags") or ""
    if _has_any_flag(flags, _MUST_CHECK_FLAGS):
        return True
    if row.get("slug") in _NBK_MUSCAT_SLUGS and _has_any_flag(flags, _NBK_MUSCAT_EPS):
        return True
    return False


def _is_for_info(row: dict[str, str]) -> bool:
    return _has_any_flag(row.get("flags") or "", _FOR_INFO_FLAGS)


_ILLEGAL_XML_RE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f￿￾]")


def _xl(val: object) -> object:
    if isinstance(val, str):
        return _ILLEGAL_XML_RE.sub("", val)
    return val


def main() -> int:
    with open(CSV_IN, newline="", encoding="utf-8") as fh:
        reader = csv.DictReader(fh)
        all_rows = list(reader)

    # Exclude pure FS_BOUNDARY_NOT_FOUND rows (no FS in PDF — separate checklist)
    flagged = [
        r for r in all_rows
        if r.get("flags") and not _is_true_no_fs(r["flags"])
    ]

    def _sort_key(r: dict[str, str]) -> tuple[str, str]:
        return (r["slug"], r["fiscal_year"])

    must_check_rows = sorted([r for r in flagged if _is_must_check(r)], key=_sort_key)
    for_info_rows   = sorted(
        [r for r in flagged if not _is_must_check(r) and _is_for_info(r)], key=_sort_key
    )

    print(f"Total rows in CSV     : {len(all_rows)}")
    print(f"Flagged (non-no-FS)   : {len(flagged)}")
    print(f"Must check            : {len(must_check_rows)}")
    print(f"For information       : {len(for_info_rows)}")

    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = openpyxl.Workbook()
    ws_must = wb.active
    ws_must.title = "Must check"
    ws_info = wb.create_sheet("For information")

    # ---- shared column layout -----------------------------------------------
    cols: list[str] = []
    cols += ["slug", "fiscal_year", "is_bank", "currency", "unit", "flag_types", "flags"]
    for f in _EV_FIELDS:
        cols += [f"{f}_value", f"{f}_page", f"{f}_label", f"{f}_confidence"]
    cols += _RATIO_FIELDS

    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(bold=True, color="FFFFFF")

    FLAG_FILL = {
        "BS_MISMATCH":              PatternFill("solid", fgColor="FFD700"),
        "NP_ATTR_GT_TOTAL":         PatternFill("solid", fgColor="FFA500"),
        "EPS_MISMATCH":             PatternFill("solid", fgColor="FF6B6B"),
        "ROE_NPA_FALLBACK":         PatternFill("solid", fgColor="C8E6C9"),
        "LEVERAGE_NEGATIVE_EQUITY": PatternFill("solid", fgColor="FF9999"),
        "FS_BOUNDARY_NOT_FOUND":    PatternFill("solid", fgColor="E0E0E0"),
        "OCR_HEAVY":                PatternFill("solid", fgColor="B0BEC5"),
    }
    ft_col_idx = cols.index("flag_types") + 1

    def _num(raw: str) -> object:
        if not raw:
            return None
        try:
            return float(raw)
        except ValueError:
            return raw

    def _write_sheet(ws: object, rows: list[dict[str, str]]) -> None:
        # Header row
        for col_i, col_name in enumerate(cols, start=1):
            cell = ws.cell(row=1, column=col_i, value=col_name)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")

        # Data rows
        for row_i, r in enumerate(rows, start=2):
            ft = _flag_types(r.get("flags", ""))
            row_vals: dict[str, object] = {
                "slug":        r["slug"],
                "fiscal_year": _num(r["fiscal_year"]),
                "is_bank":     r.get("is_bank"),
                "currency":    r.get("currency"),
                "unit":        r.get("unit"),
                "flag_types":  ft,
                "flags":       r.get("flags"),
            }
            for f in _EV_FIELDS:
                row_vals[f"{f}_value"]      = _num(r.get(f"{f}_value", ""))
                row_vals[f"{f}_page"]       = _num(r.get(f"{f}_page", ""))
                row_vals[f"{f}_label"]      = r.get(f"{f}_label")
                row_vals[f"{f}_confidence"] = r.get(f"{f}_confidence")
            for rf in _RATIO_FIELDS:
                row_vals[rf] = _num(r.get(rf, ""))

            for col_i, col_name in enumerate(cols, start=1):
                ws.cell(row=row_i, column=col_i, value=_xl(row_vals.get(col_name)))

            # Colour flag_types cell by primary flag
            ft_cell = ws.cell(row=row_i, column=ft_col_idx)
            for flag_key, fill in FLAG_FILL.items():
                if flag_key in ft:
                    ft_cell.fill = fill
                    break

        # Column widths
        for col_i, col_name in enumerate(cols, start=1):
            if "_label" in col_name:
                ws.column_dimensions[openpyxl.utils.get_column_letter(col_i)].width = 40
            elif col_name in ("slug", "flags"):
                ws.column_dimensions[openpyxl.utils.get_column_letter(col_i)].width = 55
            elif col_name == "flag_types":
                ws.column_dimensions[openpyxl.utils.get_column_letter(col_i)].width = 35
            else:
                ws.column_dimensions[openpyxl.utils.get_column_letter(col_i)].width = max(12, len(col_name) + 2)

        ws.freeze_panes = "A2"

    _write_sheet(ws_must, must_check_rows)
    _write_sheet(ws_info, for_info_rows)

    wb.save(XLSX_OUT)
    print(f"XLSX -> {XLSX_OUT}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
