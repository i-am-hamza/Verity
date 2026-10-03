"""Load institutions from data/List of Companies.xlsx.

Runs from `backend/`:
    python -m seed.load_institutions

Behaviour:
- Idempotent. Slug (kebab-case of Name) is the natural key; rerunning
  updates fields on existing rows rather than duplicating them.
- Reads industry EXACTLY as written in the file. Does not reclassify.
- Sets is_financial = (industry == "Financial Services") from the file.
- Sets wave = "financial" if is_financial else "other".
- Leaves institution_type NULL unless data/institution_types_proposed.csv
  has a non-empty "confirmed" column for that slug.
- Regenerates data/institution_types_proposed.csv on each run: proposals
  come from name heuristics but are written to a "proposed_type" column,
  never to the "confirmed" column, which is the human's responsibility.
"""
from __future__ import annotations

import csv
import re
import unicodedata
from datetime import date
from pathlib import Path

from openpyxl import load_workbook

from app.database import SessionLocal
from app.models import (  # noqa: F401 — force mapper registration
    institution as _institution,
    provenance as _provenance,
    report as _report,
    score as _score,
    taxonomy as _taxonomy,
)
from app.models.institution import Institution, InstitutionType, Wave

REPO_ROOT = Path(__file__).resolve().parents[2]
XLSX_PATH = REPO_ROOT / "data" / "List of Companies.xlsx"
PROPOSED_CSV = REPO_ROOT / "data" / "institution_types_proposed.csv"
MARKET_CAP_DATE = date(2024, 4, 16)


def slugify(name: str) -> str:
    """Kebab-case slug. ASCII-only; punctuation stripped; single hyphens."""
    normalised = unicodedata.normalize("NFKD", name)
    ascii_only = normalised.encode("ascii", "ignore").decode("ascii")
    lowered = ascii_only.lower()
    hyphenated = re.sub(r"[^a-z0-9]+", "-", lowered).strip("-")
    return hyphenated or "unknown"


def read_companies(path: Path) -> list[dict[str, object]]:
    """Header is on row 2. Data rows follow. Returns a list of dicts keyed
    by the canonical column names, NOT the sheet's exact case."""
    wb = load_workbook(path, data_only=True)
    ws = wb.active
    if ws is None:
        raise RuntimeError("No active sheet")
    header = [ws.cell(row=2, column=c).value for c in range(1, ws.max_column + 1)]
    if header[:6] != ["Rank", "Name", "Symbol", "Country", "Market CAP April 16, 2024", "Industry"]:
        raise RuntimeError(f"Unexpected header on row 2: {header}")

    rows: list[dict[str, object]] = []
    for r in range(3, ws.max_row + 1):
        rank = ws.cell(row=r, column=1).value
        name = ws.cell(row=r, column=2).value
        if rank is None and name is None:
            continue
        rows.append({
            "rank": int(rank) if rank is not None else None,
            "name": str(name).strip() if name is not None else "",
            "ticker": str(ws.cell(row=r, column=3).value or "").strip() or None,
            "country": str(ws.cell(row=r, column=4).value or "").strip(),
            "market_cap_usd": int(ws.cell(row=r, column=5).value)
                if ws.cell(row=r, column=5).value is not None else None,
            "industry": str(ws.cell(row=r, column=6).value or "").strip() or None,
        })
    return rows


# Exclusion hints for institution_type — used ONLY to peel non-bank
# financials off. Anything financial that doesn't match any of these
# defaults to `bank`, per the original spec ("rest of Financial Services
# rows -> bank"). Patch B replaces the previous "keyword-hunt for 'bank'"
# heuristic that made Kuwait Finance House, Emirates NBD and QNB depend
# on incidental spelling quirks.
_NON_BANK_HINTS: list[tuple[re.Pattern[str], InstitutionType]] = [
    (re.compile(r"\bbupa\b|\binsurance\b|\btakaful\b", re.I), InstitutionType.insurer),
    (re.compile(r"\bleasing\b|\bALAFCO\b|aircraft.*leasing", re.I), InstitutionType.leasing),
    (re.compile(
        r"investment|holding|industrial investment|development.*investment|dubai investment",
        re.I,
    ), InstitutionType.investment_holding),
]


def propose_type(name: str, industry: str | None) -> InstitutionType | None:
    """Return a *proposal* (not a decision).

    Rules:
    - Non-financial row -> `other` (the enum's catch-all).
    - Financial row that matches any non-bank hint -> that specific type.
    - Financial row that matches nothing -> `bank` (the deliberate default
      for the "Financial Services" bucket, since Middle East banks come
      in many name forms — "Bank X", "X Finance House", "X NBD", "QNB" —
      and the reliable signal is exclusion, not inclusion).
    """
    if industry != "Financial Services":
        return InstitutionType.other
    for pat, t in _NON_BANK_HINTS:
        if pat.search(name):
            return t
    return InstitutionType.bank


def read_confirmed_types(csv_path: Path) -> dict[str, InstitutionType]:
    """Read only rows where the human filled in a valid `confirmed` value.
    Silently ignores unknown values (typed but not recognised)."""
    if not csv_path.exists():
        return {}
    confirmed: dict[str, InstitutionType] = {}
    valid = {t.value: t for t in InstitutionType}
    with csv_path.open("r", encoding="utf-8", newline="") as fh:
        reader = csv.DictReader(fh)
        for row in reader:
            slug = (row.get("slug") or "").strip()
            value = (row.get("confirmed") or "").strip().lower()
            if slug and value in valid:
                confirmed[slug] = valid[value]
    return confirmed


def write_proposed_csv(csv_path: Path, rows: list[dict[str, object]]) -> None:
    """Regenerates the proposals file. Preserves existing `confirmed` values."""
    existing_confirmed: dict[str, str] = {}
    if csv_path.exists():
        with csv_path.open("r", encoding="utf-8", newline="") as fh:
            for row in csv.DictReader(fh):
                slug = (row.get("slug") or "").strip()
                if slug:
                    existing_confirmed[slug] = (row.get("confirmed") or "").strip()

    csv_path.parent.mkdir(parents=True, exist_ok=True)
    with csv_path.open("w", encoding="utf-8", newline="") as fh:
        writer = csv.writer(fh)
        writer.writerow(
            ["slug", "name", "country", "industry", "is_financial", "proposed_type", "confirmed"]
        )
        for row in rows:
            slug = slugify(str(row["name"]))
            industry = row["industry"]
            is_financial = industry == "Financial Services"
            proposed = propose_type(str(row["name"]), industry if isinstance(industry, str) else None)
            writer.writerow([
                slug,
                row["name"],
                row["country"],
                industry or "",
                "true" if is_financial else "false",
                proposed.value if proposed else "",
                existing_confirmed.get(slug, ""),
            ])


def load(dry_run: bool = False) -> dict[str, int]:
    rows = read_companies(XLSX_PATH)
    write_proposed_csv(PROPOSED_CSV, rows)
    confirmed_types = read_confirmed_types(PROPOSED_CSV)

    # Precompute slugs and detect any collisions (two rows sharing a slug).
    slug_counts: dict[str, int] = {}
    for row in rows:
        slug_counts[slugify(str(row["name"]))] = slug_counts.get(slugify(str(row["name"])), 0) + 1
    duplicate_slugs = [s for s, n in slug_counts.items() if n > 1]
    if duplicate_slugs:
        raise RuntimeError(
            f"slug collision in the company list: {duplicate_slugs} — fix the name or slugify()"
        )

    counts_by_country: dict[str, int] = {}
    counts_by_industry: dict[str, int] = {}
    counts_by_wave: dict[str, int] = {}

    db = SessionLocal()
    try:
        for row in rows:
            slug = slugify(str(row["name"]))
            industry = row["industry"] if isinstance(row["industry"], str) else None
            is_financial = industry == "Financial Services"
            wave = Wave.financial if is_financial else Wave.other

            counts_by_country[str(row["country"])] = counts_by_country.get(str(row["country"]), 0) + 1
            counts_by_industry[industry or "(missing)"] = (
                counts_by_industry.get(industry or "(missing)", 0) + 1
            )
            counts_by_wave[wave.value] = counts_by_wave.get(wave.value, 0) + 1

            existing = db.query(Institution).filter(Institution.slug == slug).first()
            fields = {
                "name": row["name"],
                "country": row["country"],
                "rank": row["rank"],
                "ticker": row["ticker"],
                "market_cap_usd": row["market_cap_usd"],
                "market_cap_date": MARKET_CAP_DATE,
                "industry": industry,
                "slug": slug,
                "is_financial": is_financial,
                "wave": wave,
                "institution_type": confirmed_types.get(slug),
            }
            if existing:
                for k, v in fields.items():
                    setattr(existing, k, v)
            else:
                db.add(Institution(**fields))

        if dry_run:
            db.rollback()
        else:
            db.commit()
    finally:
        db.close()

    return {
        "total_rows": len(rows),
        "financial_rows": sum(1 for r in rows if r["industry"] == "Financial Services"),
        "counts_by_country": counts_by_country,
        "counts_by_industry": counts_by_industry,
        "counts_by_wave": counts_by_wave,
        "confirmed_types_applied": len(confirmed_types),
    }


def main() -> None:
    summary = load()
    print(f"Loaded {summary['total_rows']} institutions ({summary['financial_rows']} financial).")
    print(f"Confirmed institution_type values applied: {summary['confirmed_types_applied']}")
    print()
    print("By country:")
    for country, n in sorted(summary["counts_by_country"].items(), key=lambda x: (-x[1], x[0])):
        print(f"  {country:>25} {n:>3}")
    print()
    print("By industry:")
    for industry, n in sorted(summary["counts_by_industry"].items(), key=lambda x: (-x[1], x[0])):
        print(f"  {industry:>25} {n:>3}")
    print()
    print("By wave:")
    for wave, n in sorted(summary["counts_by_wave"].items()):
        print(f"  {wave:>25} {n:>3}")


if __name__ == "__main__":
    main()
